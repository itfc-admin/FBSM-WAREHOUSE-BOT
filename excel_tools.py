from __future__ import annotations

import io
import re
from collections import defaultdict
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

CATEGORY_ALIASES = {
    "T-SHIRT": ["TSHIRT", "T SHIRT", "T-SHIRT", "TEE", "TEE SHIRT"],
    "LONGSLEEVE": ["LONGSLEEVE", "LONG SLEEVE", "LONG-SLEEVE", "LS TSHIRT", "L/S TSHIRT"],
    "SHIRT": ["SHIRT"],
    "DENIM SHIRT": ["DENIM SHIRT", "JEAN SHIRT", "JEANS SHIRT"],
    "SWEATSHIRT": ["SWEATSHIRT", "SWEAT SHIRT", "SWEAT-SHIRT"],
    "HOODIE": ["HOODIE", "HOODY", "HOODED SWEATSHIRT"],
    "SWEATER": ["SWEATER", "PULLOVER"],
    "KNIT": ["KNIT", "KNITWEAR", "KNITTED"],
    "PANTS": ["PANTS", "TROUSERS", "TROUSER"],
    "TRAINING PANTS": ["TRAINING PANTS", "TRACK PANTS", "JOGGER", "JOGGERS", "SWEAT PANTS", "SWEATPANTS"],
    "DENIM PANTS": ["DENIM PANTS", "DENIM PT", "JEANS", "JEAN PANTS", "DENIM TROUSERS"],
    "SHORTS": ["SHORT", "SHORTS", "SHORT PANTS"],
    "SKIRT": ["SKIRT", "SKIRTS"],
    "DRESS": ["DRESS", "DRESSES"],
    "JACKET": ["JACKET", "JACKETS"],
    "COAT": ["COAT", "COATS"],
    "VEST": ["VEST", "VESTS", "GILET"],
}

HEADER_HINTS = {
    "category": ["CATEGORY", "ITEM", "ITEM TYPE", "PRODUCT", "PRODUCT TYPE", "DESCRIPTION", "TYPE", "품목", "카테고리"],
    "qty": ["QTY", "QUANTITY", "PCS", "EA", "COUNT", "수량", "QUANTITY(EA)"],
    "brand": ["BRAND", "브랜드"],
}


def _clean(v) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).strip()).upper()


def normalize_category(value: str) -> str:
    x = _clean(value).replace("_", " ")
    x = re.sub(r"\s+", " ", x)
    for canonical, aliases in CATEGORY_ALIASES.items():
        for alias in aliases:
            a = _clean(alias)
            if x == a or (len(a) >= 5 and a in x):
                return canonical
    return "OTHER"


def _find_header_row(ws, scan_rows=20):
    best = None
    for r in range(1, min(ws.max_row, scan_rows) + 1):
        values = [_clean(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        found = {}
        for kind, hints in HEADER_HINTS.items():
            for c, value in enumerate(values, 1):
                if any(value == h or h in value for h in hints):
                    found[kind] = c
                    break
        score = int("category" in found) + int("qty" in found) + int("brand" in found)
        if score >= 2 and (best is None or score > best[0]):
            best = (score, r, found)
    return best


def parse_packing_list(data: bytes, fallback_brand: str | None = None):
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    totals = defaultdict(int)
    detected_brands = defaultdict(int)
    parsed_rows = 0
    warnings = []

    for ws in wb.worksheets:
        header = _find_header_row(ws)
        if not header:
            warnings.append(f"Лист '{ws.title}': не найдены колонки Category/Qty")
            continue
        _, header_row, cols = header
        for r in range(header_row + 1, ws.max_row + 1):
            raw_category = ws.cell(r, cols["category"]).value if "category" in cols else None
            raw_qty = ws.cell(r, cols["qty"]).value if "qty" in cols else None
            raw_brand = ws.cell(r, cols["brand"]).value if "brand" in cols else None
            if raw_category in (None, "") and raw_qty in (None, ""):
                continue
            try:
                qty = int(round(float(str(raw_qty).replace(",", "").strip())))
            except Exception:
                continue
            if qty <= 0:
                continue
            cat = normalize_category(str(raw_category or "OTHER"))
            totals[cat] += qty
            parsed_rows += 1
            if raw_brand:
                detected_brands[_clean(raw_brand)] += qty

    if not totals:
        raise ValueError("В Excel не удалось найти строки товара с Category/Qty.")

    brand = _clean(fallback_brand or "")
    if not brand and detected_brands:
        brand = max(detected_brands.items(), key=lambda x: x[1])[0]
    if not brand:
        raise ValueError("Не удалось определить бренд. Укажите бренд перед загрузкой файла.")

    return {
        "brand": brand,
        "items": dict(sorted(totals.items())),
        "total_qty": sum(totals.values()),
        "parsed_rows": parsed_rows,
        "warnings": warnings,
    }


def create_order_packing_xlsx(order_no: str, client_code: str, brand: str, items: list[tuple[str, int]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Packing List"

    ws["A1"] = "FBSM WAREHOUSE — PACKING LIST"
    ws["A1"].font = Font(size=16, bold=True)
    ws.merge_cells("A1:C1")
    ws["A3"] = "Order"
    ws["B3"] = order_no
    ws["A4"] = "Client"
    ws["B4"] = client_code
    ws["A5"] = "Brand"
    ws["B5"] = brand

    headers = ["Brand", "Category", "Qty (EA)"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(7, c, h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center")

    thin = Side(style="thin", color="B7B7B7")
    total = 0
    for i, (category, qty) in enumerate(items, 8):
        ws.cell(i, 1, brand)
        ws.cell(i, 2, category)
        ws.cell(i, 3, qty)
        total += qty
        for c in range(1, 4):
            ws.cell(i, c).border = Border(bottom=thin)

    end = 7 + len(items) + 1
    ws.cell(end, 2, "TOTAL").font = Font(bold=True)
    ws.cell(end, 3, total).font = Font(bold=True)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 14
    ws.freeze_panes = "A8"

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
