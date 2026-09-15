from pathlib import Path
from collections import defaultdict
import json
from utils import clean, classify_category

def _rows_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    result = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([v for v in row])
        result.append((ws.title, rows))
    return result

def _rows_xls(path):
    import xlrd
    wb = xlrd.open_workbook(path)
    result = []
    for ws in wb.sheets():
        rows = []
        for r in range(ws.nrows):
            rows.append([ws.cell_value(r,c) for c in range(ws.ncols)])
        result.append((ws.name, rows))
    return result

def read_workbook(path):
    suffix = Path(path).suffix.lower()
    if suffix == ".xlsx":
        return _rows_xlsx(path)
    if suffix == ".xls":
        return _rows_xls(path)
    raise ValueError("Поддерживаются только .xls и .xlsx")

def norm(v):
    return clean(v).upper().replace("\n"," ").strip()

def find_header(rows, required_groups, max_scan=40):
    for idx, row in enumerate(rows[:max_scan]):
        vals = [norm(v) for v in row]
        ok = True
        found = {}
        for key, aliases in required_groups.items():
            pos = next((i for i,v in enumerate(vals) if any(a in v for a in aliases)), None)
            if pos is None:
                ok = False
                break
            found[key] = pos
        if ok:
            return idx, found, vals
    return None, None, None

def detect_template(rows):
    # CARIES NOTE: 품번 / 칼라 / 사이즈 / 수량, often #BOX.NO
    idx, cols, vals = find_header(rows, {
        "style":["품번"],
        "color":["칼라","COLOR","COL"],
        "size":["사이즈","SIZE"],
        "qty":["수량","QTY","QUANTITY"],
    })
    if idx is not None:
        box = next((i for i,v in enumerate(vals) if "BOX" in v), None)
        return "CARIES_BOX", idx, {**cols, "box":box}

    # TOFFEE: 스타일넘버 / 상품명 / 총수량, sizes in separate columns
    idx, cols, vals = find_header(rows, {
        "style":["스타일넘버","STYLE NUMBER","STYLE NO"],
        "product":["상품명","PRODUCT NAME"],
        "total":["총수량","총 수량","TOTAL QTY","TOTAL"],
    })
    if idx is not None:
        size_cols = {}
        for s in ["XXS","XS","S","M","L","XL","XXL","2XL","3XL","FREE"]:
            p = next((i for i,v in enumerate(vals) if v == s), None)
            if p is not None:
                size_cols[s] = p
        return "TOFFEE_PRODUCT", idx, {**cols, "size_cols":size_cols}

    # PLAC: 품목명 / STYLE / SIZE / 내품수량 / SKU
    idx, cols, vals = find_header(rows, {
        "product":["품목명","PRODUCT NAME"],
        "style":["STYLE"],
        "size":["SIZE"],
        "qty":["내품수량","수량","QTY"],
        "sku":["SKU"],
    })
    if idx is not None:
        color = next((i for i,v in enumerate(vals) if v in ("COL","COLOR","칼라")), None)
        box = next((i for i,v in enumerate(vals) if "BOX" in v or "박스" in v), None)
        return "PLAC_PACKING", idx, {**cols, "color":color, "box":box}

    # Generic CATEGORY / QTY
    idx, cols, vals = find_header(rows, {
        "category":["CATEGORY","CAT","TYPE"],
        "qty":["QTY","QUANTITY","EA","PCS"],
    })
    if idx is not None:
        brand = next((i for i,v in enumerate(vals) if v in ("BRAND","BRAND NAME")), None)
        return "GENERIC_CATEGORY", idx, {**cols, "brand":brand}

    return None, None, None

def safe_int(v):
    if v is None or clean(v)=="":
        return 0
    try:
        return int(round(float(str(v).replace(",",""))))
    except Exception:
        return 0

def cell(row, idx):
    if idx is None or idx >= len(row):
        return ""
    return row[idx]

def parse_caries(rows, h, c):
    details=[]
    for row in rows[h+1:]:
        style=clean(cell(row,c["style"]))
        qty=safe_int(cell(row,c["qty"]))
        if not style or qty<=0:
            continue
        color=clean(cell(row,c["color"]))
        size=clean(cell(row,c["size"]))
        box=clean(cell(row,c.get("box")))
        cat=classify_category("",style,"")
        details.append({
            "category":cat,"box_no":box,"style_no":style,"product_name":"",
            "color":color,"size":size,"sku":"","qty":qty
        })
    return details

def parse_toffee(rows, h, c):
    details=[]
    size_cols=c.get("size_cols",{})
    for row in rows[h+1:]:
        style=clean(cell(row,c["style"]))
        product=clean(cell(row,c["product"]))
        if not style and not product:
            continue
        total=safe_int(cell(row,c["total"]))
        cat=classify_category(product,style,"")
        size_sum=0
        for size,idx in size_cols.items():
            q=safe_int(cell(row,idx))
            if q>0:
                details.append({
                    "category":cat,"box_no":"","style_no":style,"product_name":product,
                    "color":"","size":size,"sku":"","qty":q
                })
                size_sum += q
        if not size_cols and total>0:
            details.append({
                "category":cat,"box_no":"","style_no":style,"product_name":product,
                "color":"","size":"","sku":"","qty":total
            })
        elif size_cols and total>0 and size_sum==0:
            details.append({
                "category":cat,"box_no":"","style_no":style,"product_name":product,
                "color":"","size":"","sku":"","qty":total
            })
    return details

def parse_plac(rows, h, c):
    details=[]
    for row in rows[h+1:]:
        product=clean(cell(row,c["product"]))
        style=clean(cell(row,c["style"]))
        sku=clean(cell(row,c["sku"]))
        qty=safe_int(cell(row,c["qty"]))
        if not (product or style or sku) or qty<=0:
            continue
        color=clean(cell(row,c.get("color")))
        size=clean(cell(row,c["size"]))
        box=clean(cell(row,c.get("box")))
        cat=classify_category(product,style,sku)
        details.append({
            "category":cat,"box_no":box,"style_no":style,"product_name":product,
            "color":color,"size":size,"sku":sku,"qty":qty
        })
    return details

def parse_generic(rows,h,c):
    details=[]
    for row in rows[h+1:]:
        cat=clean(cell(row,c["category"])).upper()
        qty=safe_int(cell(row,c["qty"]))
        if not cat or qty<=0:
            continue
        details.append({
            "category":cat,"box_no":"","style_no":"","product_name":"",
            "color":"","size":"","sku":"","qty":qty
        })
    return details

def parse_universal(path):
    sheets=read_workbook(path)
    best=None
    for sheet_name, rows in sheets:
        fmt,h,c=detect_template(rows)
        if not fmt:
            continue
        if fmt=="CARIES_BOX":
            details=parse_caries(rows,h,c)
        elif fmt=="TOFFEE_PRODUCT":
            details=parse_toffee(rows,h,c)
        elif fmt=="PLAC_PACKING":
            details=parse_plac(rows,h,c)
        else:
            details=parse_generic(rows,h,c)
        if details:
            qty=sum(x["qty"] for x in details)
            if best is None or qty>best["total"]:
                best={"format":fmt,"sheet":sheet_name,"details":details,"total":qty}
    if not best:
        headers=[]
        for name,rows in sheets[:3]:
            for row in rows[:10]:
                vals=[clean(v) for v in row if clean(v)]
                if vals:
                    headers.append(f"{name}: {' | '.join(vals[:12])}")
        raise ValueError("Неизвестный формат Excel. Найденные строки:\n" + "\n".join(headers[:8]))

    cats=defaultdict(int)
    styles=set()
    boxes=set()
    unknown=set()
    for d in best["details"]:
        cats[d["category"]]+=d["qty"]
        if d["style_no"]: styles.add(d["style_no"])
        if d["box_no"]: boxes.add(d["box_no"])
        if d["category"]=="OTHER":
            unknown.add(d["style_no"] or d["sku"] or d["product_name"] or "?")
    best["categories"]=dict(sorted(cats.items()))
    best["styles_count"]=len(styles)
    best["boxes_count"]=len(boxes)
    best["unknown"]=sorted(unknown)[:30]
    return best
